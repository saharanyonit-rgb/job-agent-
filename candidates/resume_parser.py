from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any


class ResumeParseError(Exception):
    pass


@dataclass
class ResumeText:
    text: str
    file_type: str
    filename: str
    content_hash: str


@dataclass
class ResumeParsed:
    summary: str | None = None
    skills: list[str] = None  # type: ignore[assignment]
    technologies: list[str] = None  # type: ignore[assignment]
    years_of_experience: float | None = None
    certifications: list[str] = None  # type: ignore[assignment]
    education: list[dict[str, Any]] = None  # type: ignore[assignment]
    projects: list[str] = None  # type: ignore[assignment]
    previous_employers: list[str] = None  # type: ignore[assignment]
    current_title: str | None = None
    achievements: list[str] = None  # type: ignore[assignment]

    def to_dict(self) -> dict[str, Any]:
        return {
            "summary": self.summary,
            "skills": self.skills or [],
            "technologies": self.technologies or [],
            "years_of_experience": self.years_of_experience,
            "certifications": self.certifications or [],
            "education": self.education or [],
            "projects": self.projects or [],
            "previous_employers": self.previous_employers or [],
            "current_title": self.current_title,
            "achievements": self.achievements or [],
        }


def extract_text(path: str | Path) -> ResumeText:
    path = Path(path)
    if not path.exists():
        raise ResumeParseError(f"File not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(path)
        file_type = "pdf"
    elif suffix in (".docx", ".doc"):  # .doc legacy not supported
        text = _extract_docx(path)
        file_type = "docx"
    elif suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="replace")
        file_type = "txt"
    else:
        raise ResumeParseError(f"Unsupported file type: {suffix} (supported: pdf, docx, txt)")

    import hashlib

    content_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    return ResumeText(text=text, file_type=file_type, filename=path.name, content_hash=content_hash)


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


_SKILL_LEXICON = [
    "python", "java", "javascript", "typescript", "c#", ".net", "go", "rust", "c++",
    "sql", "postgresql", "mysql", "mongodb", "redis", "kafka", "rabbitmq",
    "azure", "aws", "gcp", "kubernetes", "docker", "terraform", "ansible",
    "ci/cd", "gitlab ci", "jenkins", "github actions", "devops", "sre",
    "microservices", "rest api", "graphql", "grpc", "event-driven",
    "react", "angular", "vue", "node.js", "express", "next.js", "fastapi",
    "django", "flask", "spring boot", "spring", "hibernate", "fastapi",
    "llm", "rag", "generative ai", "genai", "langchain", "pytorch",
    "tensorflow", "mlops", "airflow", "spark", "hadoop", "databricks", "dbt",
    "architecture", "solution architecture", "technical leadership",
    "tableau", "power bi", "analytics", "etl", "data engineering",
    "oauth", "jwt", "saml", "security", "cybersecurity",
    "scrum", "agile", "kanban", "project management",
]


class ResumeParser:
    def parse(self, resume_text: ResumeText) -> ResumeParsed:
        text = resume_text.text
        lowered = text.lower()
        langs = {"python", "java", "javascript", "typescript", "go", "rust", "sql", "postgresql", "mysql", "mongodb", "redis", "kafka", "azure", "aws", "gcp", "kubernetes", "docker", "terraform", "react", "angular", "vue", "node.js", "fastapi", "django", "flask", "spring", "pytorch", "tensorflow", "spark", "airflow"}
        tech = [s for s in _SKILL_LEXICON if _contains_token(lowered, s)]
        parsed = ResumeParsed(
            skills=tech,
            technologies=[t for t in tech if t in langs],
            projects=[line.strip() for line in text.splitlines() if line.strip().lower().startswith("project")],
            achievements=[line.strip() for line in text.splitlines() if _looks_like_achievement(line)],
            certifications=_find_certifications(text),
            education=_find_education(text),
            previous_employers=_find_employers(text),
            years_of_experience=_find_years_experience(text),
            current_title=_find_titles(text),
            summary=_find_summary(text),
        )
        return parsed


def _contains_token(text: str, token: str) -> bool:
    import re

    return bool(re.search(rf"(?<![a-z0-9]){re.escape(token)}(?![a-z0-9])", text))


def _find_years_experience(text: str) -> float | None:
    import re

    m = re.search(r"(\d{1,2})(?:\.\d)?\s*(?:\+)?\s*(?:years?|yrs)[^a-z]", text, re.I)
    if m:
        return float(m.group(1))
    m = re.search(r"experience (?:of )?(\d{1,2})\s*(?:years?|yrs)", text, re.I)
    return float(m.group(1)) if m else None


def _find_certifications(text: str) -> list[str]:
    import re

    patterns = [
        r"(?:certified|certification|certificate)[^.\n]{0,80}",
        r"\b(?:AWS|AZ|Google|Microsoft|PMP|TOGAF)\b[^.\n]{0,40}",
        r"\b(?:Solution Architect|Developer Associate|Solutions Architect)\b",
    ]
    out: list[str] = []
    for pat in patterns:
        for m in re.finditer(pat, text, re.I):
            val = m.group(0).strip()
            if val and val.lower() not in [o.lower() for o in out]:
                out.append(val)
    return out[:12]


def _find_education(text: str) -> list[dict]:
    import re

    out: list[dict] = []
    for m in re.finditer(r"(B\.?(?:Tech|E|A|Sc)|M\.?(?:Tech|BA|Sc)|MBA)", text):
        val = m.group(1)
        if val not in [e["degree"] for e in out]:
            out.append({"degree": val})
    return out


def _find_employers(text: str) -> list[str]:
    import re

    # Lines under a "Experience" header that look like company names (Capitalized words)
    lines = text.splitlines()
    companies: list[str] = []
    in_experience = False
    markers = ("experience", "employment history", "work history")
    for line in lines:
        stripped = line.strip()
        low = stripped.lower()
        if any(m in low for m in markers):
            in_experience = True
            continue
        if in_experience:
            if re.match(r"^[A-Z][A-Za-z0-9&.\- ]{2,60}$", stripped) and len(stripped.split()) <= 6:
                if stripped.lower() not in [c.lower() for c in companies]:
                    companies.append(stripped)
            if re.match(r"^(20|19)\d\d\s*[-–]\s*(20|19)?\d\d|present|current", low):
                continue
            if len(companies) >= 8:
                break
    return companies


_TITLES = [
    "Technical Architect", "Solution Architect", "Cloud Architect", "AI Architect",
    "Software Engineer", "Senior Software Engineer", "Lead Engineer", "Principal Engineer",
    "Engineering Manager", "Data Engineer", "Devops Engineer", "Data Scientist",
    "Product Manager", "Full Stack Developer", "Backend Engineer",
]


def _find_titles(text: str) -> str | None:
    lowered = text.lower()
    for title in _TITLES:
        if title.lower() in lowered:
            return title
    return None


def _looks_like_achievement(line: str) -> bool:
    import re

    return bool(re.search(r"\d+\s*%|\$\s?[\d,]+k?|\d+\s*(?:users?|customers?|requests)/", line, re.I)) and len(line) < 200


def _find_summary(text: str) -> str | None:
    import re

    m = re.search(r"(?:summary|profile|about)\s*:\s*([^\n]{20,300})", text, re.I)
    return m.group(1).strip() if m else None